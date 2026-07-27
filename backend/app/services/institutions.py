"""US-INS-02 — Institution CRUD and Excel import."""

from __future__ import annotations

import io
import re
import uuid
from typing import Any

from fastapi import status
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, or_, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.errors import AppError, FieldError
from app.models import Institution, Region, User
from app.services.audit import write_audit_event
from app.services.storage import get_object_storage

_CODE_RE = re.compile(r"^[A-Za-z0-9][A-Za-z0-9._\-]{0,63}$")

_HEADER_ALIASES = {
    "school code": "code",
    "code": "code",
    "name": "name",
    "school name": "name",
    "region": "region",
}


def _institution_out(inst: Institution) -> dict[str, Any]:
    return {
        "institutionId": str(inst.id),
        "code": inst.code,
        "name": inst.name,
        "regionId": str(inst.region_id) if inst.region_id else None,
        "active": inst.active,
    }


async def list_institutions(
    session: AsyncSession,
    *,
    q: str | None = None,
    active: bool | None = True,
    region_id: uuid.UUID | None = None,
) -> list[dict[str, Any]]:
    stmt = select(Institution).order_by(Institution.name)
    if active is not None:
        stmt = stmt.where(Institution.active.is_(active))
    if region_id is not None:
        stmt = stmt.where(Institution.region_id == region_id)
    if q:
        like = f"%{q.strip()}%"
        stmt = stmt.where(
            or_(Institution.name.ilike(like), Institution.code.ilike(like))
        )
    rows = (await session.execute(stmt)).scalars().all()
    return [_institution_out(i) for i in rows]


async def search_active(
    session: AsyncSession,
    *,
    q: str,
    limit: int = 25,
) -> list[dict[str, str]]:
    """Public search of active schools by name or code (registration UX)."""
    cleaned = q.strip()
    if len(cleaned) < 1:
        return []
    like = f"%{cleaned}%"
    stmt = (
        select(Institution)
        .where(
            Institution.active.is_(True),
            or_(Institution.name.ilike(like), Institution.code.ilike(like)),
        )
        .order_by(Institution.name)
        .limit(max(1, min(limit, 50)))
    )
    rows = (await session.execute(stmt)).scalars().all()
    return [
        {"code": i.code, "name": i.name, "institutionId": str(i.id)} for i in rows
    ]


async def lookup_by_code(session: AsyncSession, code: str) -> dict[str, str]:
    cleaned = code.strip()
    if not cleaned:
        raise AppError(
            "SCHOOL_NOT_FOUND",
            "School not found",
            status_code=status.HTTP_404_NOT_FOUND,
            fields=[FieldError("code", "SCHOOL_NOT_FOUND")],
        )
    inst = (
        await session.execute(
            select(Institution).where(func.lower(Institution.code) == cleaned.lower())
        )
    ).scalar_one_or_none()
    if inst is None:
        raise AppError(
            "SCHOOL_NOT_FOUND",
            "School not found",
            status_code=status.HTTP_404_NOT_FOUND,
            fields=[FieldError("code", "SCHOOL_NOT_FOUND")],
        )
    if not inst.active:
        raise AppError(
            "SCHOOL_INACTIVE",
            "School is inactive",
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            fields=[FieldError("code", "SCHOOL_INACTIVE")],
        )
    return {"code": inst.code, "name": inst.name, "institutionId": str(inst.id)}


async def get_active_by_code(session: AsyncSession, code: str) -> Institution:
    cleaned = code.strip()
    if not cleaned:
        raise AppError(
            "SCHOOL_NOT_FOUND",
            "School not found",
            status_code=status.HTTP_404_NOT_FOUND,
            fields=[FieldError("schoolCode", "SCHOOL_NOT_FOUND")],
        )
    inst = (
        await session.execute(
            select(Institution).where(func.lower(Institution.code) == cleaned.lower())
        )
    ).scalar_one_or_none()
    if inst is None:
        raise AppError(
            "SCHOOL_NOT_FOUND",
            "School not found",
            status_code=status.HTTP_404_NOT_FOUND,
            fields=[FieldError("schoolCode", "SCHOOL_NOT_FOUND")],
        )
    if not inst.active:
        raise AppError(
            "SCHOOL_INACTIVE",
            "School is inactive",
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            fields=[FieldError("schoolCode", "SCHOOL_INACTIVE")],
        )
    return inst


async def _load_region(session: AsyncSession, region_id: uuid.UUID) -> Region:
    region = (
        await session.execute(select(Region).where(Region.id == region_id))
    ).scalar_one_or_none()
    if region is None or not region.active:
        raise AppError(
            "INVALID",
            "Region not found or inactive",
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            fields=[FieldError("regionId", "INVALID")],
        )
    return region


async def _resolve_region_by_name(session: AsyncSession, name: str) -> Region | None:
    return (
        await session.execute(
            select(Region).where(
                func.lower(Region.name) == name.strip().lower(),
                Region.active.is_(True),
            )
        )
    ).scalar_one_or_none()


def _validate_code(code: str) -> str:
    cleaned = code.strip()
    if not cleaned:
        raise AppError(
            "REQUIRED",
            "School code is required",
            status_code=422,
            fields=[FieldError("code", "REQUIRED")],
        )
    if not _CODE_RE.match(cleaned):
        raise AppError(
            "INVALID",
            "School code has invalid characters",
            status_code=422,
            fields=[FieldError("code", "INVALID")],
        )
    return cleaned


async def create_institution(
    session: AsyncSession,
    *,
    actor: User,
    code: str,
    name: str,
    region_id: uuid.UUID,
    active: bool = True,
    ip: str | None = None,
    user_agent: str | None = None,
) -> dict[str, Any]:
    code = _validate_code(code)
    name = name.strip()
    if not name:
        raise AppError(
            "REQUIRED",
            "Name is required",
            status_code=422,
            fields=[FieldError("name", "REQUIRED")],
        )
    await _load_region(session, region_id)
    existing = (
        await session.execute(select(Institution).where(Institution.code == code))
    ).scalar_one_or_none()
    if existing is not None:
        raise AppError(
            "DUPLICATE",
            "School code already exists",
            status_code=409,
            fields=[FieldError("code", "DUPLICATE")],
        )
    inst = Institution(code=code, name=name, region_id=region_id, active=active)
    session.add(inst)
    await session.flush()
    await write_audit_event(
        session,
        competition_id=None,
        actor_id=actor.id,
        actor_role=actor.role.value,
        action="INSTITUTION_CREATED",
        entity_type="Institution",
        entity_id=str(inst.id),
        before=None,
        after=_institution_out(inst),
        ip=ip,
        user_agent=user_agent,
    )
    await session.commit()
    await session.refresh(inst)
    return _institution_out(inst)


async def patch_institution(
    session: AsyncSession,
    *,
    actor: User,
    institution_id: uuid.UUID,
    code: str | None = None,
    name: str | None = None,
    region_id: uuid.UUID | None = None,
    active: bool | None = None,
    ip: str | None = None,
    user_agent: str | None = None,
) -> dict[str, Any]:
    inst = (
        await session.execute(select(Institution).where(Institution.id == institution_id))
    ).scalar_one_or_none()
    if inst is None:
        raise AppError("NOT_FOUND", "Institution not found", status_code=404)
    before = _institution_out(inst)
    if code is not None:
        new_code = _validate_code(code)
        if new_code != inst.code:
            clash = (
                await session.execute(select(Institution).where(Institution.code == new_code))
            ).scalar_one_or_none()
            if clash is not None:
                raise AppError(
                    "DUPLICATE",
                    "School code already exists",
                    status_code=409,
                    fields=[FieldError("code", "DUPLICATE")],
                )
            inst.code = new_code
    if name is not None:
        cleaned = name.strip()
        if not cleaned:
            raise AppError(
                "REQUIRED",
                "Name is required",
                status_code=422,
                fields=[FieldError("name", "REQUIRED")],
            )
        inst.name = cleaned
    if region_id is not None:
        await _load_region(session, region_id)
        inst.region_id = region_id
    if active is not None:
        inst.active = active
    await session.flush()
    after = _institution_out(inst)
    await write_audit_event(
        session,
        competition_id=None,
        actor_id=actor.id,
        actor_role=actor.role.value,
        action="INSTITUTION_UPDATED",
        entity_type="Institution",
        entity_id=str(inst.id),
        before=before,
        after=after,
        ip=ip,
        user_agent=user_agent,
    )
    await session.commit()
    await session.refresh(inst)
    return after


def build_import_template() -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Schools"
    ws.append(["school code", "name", "region"])
    ws.append(["SCH-001", "Example Secondary School", "Greater Accra"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _normalize_header(cell: Any) -> str | None:
    if cell is None:
        return None
    key = str(cell).strip().lower()
    return _HEADER_ALIASES.get(key)


async def import_institutions(
    session: AsyncSession,
    *,
    actor: User,
    data: bytes,
    filename: str,
    content_type: str | None,
    ip: str | None = None,
    user_agent: str | None = None,
) -> dict[str, Any]:
    lower_name = (filename or "").lower()
    if not lower_name.endswith(".xlsx"):
        raise AppError(
            "FILE_TYPE",
            "Only .xlsx files are accepted",
            status_code=400,
            fields=[FieldError("file", "FILE_TYPE")],
        )
    # Malware scan via storage pipeline
    store = get_object_storage()
    store.put(data, prefix="imports/institutions", filename=f"{uuid.uuid4()}.xlsx")

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise AppError(
            "FILE_TYPE",
            "Could not read Excel file",
            status_code=400,
            fields=[FieldError("file", "FILE_TYPE")],
        ) from exc

    ws = wb.active
    if ws is None:
        raise AppError("FILE_TYPE", "Workbook has no sheets", status_code=400)

    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        return {"created": 0, "updated": 0, "errors": []}

    col_map: dict[str, int] = {}
    for idx, cell in enumerate(header_row):
        mapped = _normalize_header(cell)
        if mapped:
            col_map[mapped] = idx
    for required in ("code", "name", "region"):
        if required not in col_map:
            raise AppError(
                "FILE_TYPE",
                f"Missing required column: {required}",
                status_code=400,
                fields=[FieldError("file", "FILE_TYPE")],
            )

    created = 0
    updated = 0
    errors: list[dict[str, Any]] = []

    for row_num, row in enumerate(rows_iter, start=2):
        if row is None or all(c is None or str(c).strip() == "" for c in row):
            continue
        code_raw = row[col_map["code"]] if col_map["code"] < len(row) else None
        name_raw = row[col_map["name"]] if col_map["name"] < len(row) else None
        region_raw = row[col_map["region"]] if col_map["region"] < len(row) else None

        code = str(code_raw).strip() if code_raw is not None else ""
        name = str(name_raw).strip() if name_raw is not None else ""
        region_name = str(region_raw).strip() if region_raw is not None else ""

        if not code:
            errors.append({"row": row_num, "field": "code", "reason": "IMPORT_ROW_INVALID"})
            continue
        if not name:
            errors.append({"row": row_num, "field": "name", "reason": "IMPORT_ROW_INVALID"})
            continue
        if not region_name:
            errors.append({"row": row_num, "field": "region", "reason": "IMPORT_ROW_INVALID"})
            continue

        region = await _resolve_region_by_name(session, region_name)
        if region is None:
            errors.append({"row": row_num, "field": "region", "reason": "IMPORT_ROW_INVALID"})
            continue

        existing = (
            await session.execute(select(Institution).where(Institution.code == code))
        ).scalar_one_or_none()
        if existing is None:
            session.add(
                Institution(code=code, name=name, region_id=region.id, active=True)
            )
            created += 1
        else:
            existing.name = name
            existing.region_id = region.id
            updated += 1

    await session.flush()
    await write_audit_event(
        session,
        competition_id=None,
        actor_id=actor.id,
        actor_role=actor.role.value,
        action="INSTITUTION_IMPORTED",
        entity_type="Institution",
        entity_id="bulk",
        before=None,
        after={"created": created, "updated": updated, "errorCount": len(errors)},
        ip=ip,
        user_agent=user_agent,
        reason=content_type,
    )
    await session.commit()
    return {"created": created, "updated": updated, "errors": errors}
